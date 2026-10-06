#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
export_incidents_xlsx.py — incident records (+ optional statistics) → .xlsx with Bulgarian headers.

Usage:
    python3 scripts/export_incidents_xlsx.py records.json -o "/mnt/user-data/outputs/Proizshestviya_....xlsx"
    python3 scripts/export_incidents_xlsx.py records.json --stats stats.json -o out.xlsx
    python3 scripts/export_incidents_xlsx.py r2025.json r2026.json --stats stats.json -o out.xlsx
    python3 scripts/export_incidents_xlsx.py records.json --with-address -o out.xlsx

Arguments:
    records          one or more JSON files from `_shared/scripts/fetch_incidents.py -o`
                     (merged and de-duplicated by id, like incident_stats.py)
    --stats FILE     JSON written by `incident_stats.py -o FILE`; without it the
                     "Обобщение" sheet shows the totals by kind of incident
    --with-address   add the street address column (left out by default)
    -o/--output      the .xlsx to write (required)

Sheets:
    „Произшествия“ – one row per incident, oldest first: date and time, № of the
                     телефонограма, service, station, settlement, municipality, kind,
                     object, object class, cause, state (and address with --with-address).
    „Обобщение“    – period, total, injured/dead, then the groups from the stats file.

Owner names, casualty details and coordinates are never written.
Requires openpyxl.
"""
import argparse
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import incident_stats as st  # noqa: E402

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    sys.exit("MISSING_DEPENDENCY: pip install openpyxl")

HEADER_FILL = PatternFill("solid", fgColor="D9D9D9")
BOLD = Font(bold=False)   # no bold in generated documents (_shared/conventions.md)
COLUMNS = [
    ("Дата и час", 17, lambda r: (lambda t: f"{t:%d.%m.%Y %H:%M}" if t else "")(st.signal_time(r))),
    ("№ на телефонограмата", 12, lambda r: st.clean(r.get("num_event"))),
    ("Служба", 22, lambda r: st.value(r, "service")),
    ("Звено", 26, lambda r: st.value(r, "station")),
    ("Населено място", 22, lambda r: st.value(r, "settlement")),
    ("Община", 18, lambda r: st.value(r, "municipality")),
    ("Вид произшествие", 30, lambda r: st.value(r, "kind")),
    ("Обект", 34, lambda r: st.value(r, "object")),
    ("Вид обект", 34, lambda r: st.value(r, "object_class")),
    ("Причина", 32, lambda r: st.value(r, "cause")),
    ("Състояние", 18, lambda r: st.value(r, "state")),
]
ADDRESS = ("Адрес", 36, lambda r: st.clean(r.get("address")))


def write_header(ws, row, titles):
    for c, t in enumerate(titles, 1):
        cell = ws.cell(row=row, column=c, value=t)
        cell.font, cell.fill = BOLD, HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def dash(text):
    """En dash between words, as in all fire-service documents."""
    return text.replace(" - ", " – ").replace("—", "–") if isinstance(text, str) else text


def list_sheet(wb, rows, with_address):
    ws = wb.active
    ws.title = "Произшествия"
    cols = COLUMNS[:5] + [ADDRESS] + COLUMNS[5:] if with_address else COLUMNS
    write_header(ws, 1, [c[0] for c in cols])
    ordered = sorted(rows, key=lambda r: (st.clean(r.get("dat")), st.clean(r.get("id"))))
    for i, r in enumerate(ordered, 2):
        for c, (_, _, get) in enumerate(cols, 1):
            ws.cell(row=i, column=c, value=dash(get(r)) or None)
    for c, (_, width, _) in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(len(ordered) + 1, 1)}"


def fmt_date(iso):
    return f"{iso[8:10]}.{iso[5:7]}.{iso[:4]} г." if iso else "—"


def totals_sheet(wb, stats):
    ws = wb.create_sheet("Обобщение")
    p = stats.get("period") or {}
    info = [
        ("Период на данните", f"{fmt_date(p.get('first'))} – {fmt_date(p.get('last'))}"),
        ("Общо произшествия", stats.get("total", 0)),
        ("Пострадали", (stats.get("casualties") or {}).get("injured", 0)),
        ("Загинали", (stats.get("casualties") or {}).get("dead", 0)),
    ]
    if p.get("empty_months"):
        info.append(("Месеци без данни", ", ".join(p["empty_months"])))
    for i, (k, v) in enumerate(info, 1):
        ws.cell(row=i, column=1, value=k).font = BOLD
        ws.cell(row=i, column=2, value=v)
    row = len(info) + 2

    cmp_vals = stats.get("compare_values") or []
    if stats.get("total_by_compare"):
        write_header(ws, row, [st.FIELD_LABELS.get(stats.get("compare_field"), stats.get("compare_field") or "")]
                     + cmp_vals)
        ws.cell(row=row + 1, column=1, value="Общо")
        for c, v in enumerate(cmp_vals, 2):
            ws.cell(row=row + 1, column=c, value=stats["total_by_compare"].get(v, 0))
        row += 3

    groups = stats.get("groups") or []
    if groups:
        gb = stats.get("group_by") or []
        titles = [st.FIELD_LABELS.get(f, f) for f in gb] + ["Брой"] + cmp_vals + ["Пострадали", "Загинали"]
        write_header(ws, row, titles)
        for g in groups:
            row += 1
            vals = list(g["key"]) + [g["count"]] + [g.get("by_compare", {}).get(v, 0) for v in cmp_vals] \
                + [g.get("injured", 0), g.get("dead", 0)]
            for c, v in enumerate(vals, 1):
                ws.cell(row=row, column=c, value=dash(v))
        other = stats.get("other")
        if other:
            row += 1
            vals = [f"Останалите ({other['groups']})"] + [""] * (len(gb) - 1) + [other["count"]] \
                + [other.get("by_compare", {}).get(v, 0) for v in cmp_vals]
            for c, v in enumerate(vals, 1):
                ws.cell(row=row, column=c, value=v)
    for c in range(1, 12):
        ws.column_dimensions[get_column_letter(c)].width = 34 if c == 1 else 16


def main(argv=None):
    ap = argparse.ArgumentParser(description="Incident records → .xlsx (Bulgarian headers).")
    ap.add_argument("records", nargs="+", help="JSON file(s) from fetch_incidents.py -o")
    ap.add_argument("--stats", help="JSON file from incident_stats.py -o")
    ap.add_argument("--with-address", action="store_true", help="add the street address column")
    ap.add_argument("-o", "--output", required=True, help="the .xlsx file to write")
    a = ap.parse_args(argv)

    rows, dropped = st.load(a.records)
    if a.stats:
        try:
            import json
            with open(a.stats, encoding="utf-8") as fh:
                stats = json.load(fh)
        except (OSError, ValueError):
            sys.exit(f"BAD_INPUT: {a.stats} is not a stats JSON file")
    else:
        stats = st.stats(rows, ["kind"], dropped=dropped, files=len(a.records))

    wb = Workbook()
    list_sheet(wb, rows, a.with_address)
    totals_sheet(wb, stats)
    wb.save(a.output)
    print(f"OK: {a.output} ({len(rows)} records)")


if __name__ == "__main__":
    main()
