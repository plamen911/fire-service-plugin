#!/usr/bin/env python3
"""Smoke tests for build_udostoverenie.py. Run from the skill folder:
    python3 tests/test_build.py
"""
import json
import os
import subprocess
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
SKILL = os.path.dirname(HERE)

# the tests never touch the network: no key, and an invented staff list instead of the real one
os.environ["FIRE_SERVICE_NO_KEY"] = "1"
os.environ["FIRE_SERVICE_STAFF_CSV"] = os.path.abspath(os.path.join(SKILL, "..", "..", "..", "..", "tests", "fixtures", "staff.csv"))
sys.path.insert(0, os.path.join(SKILL, "scripts"))

from docx import Document  # noqa: E402
from build_udostoverenie import forces_sentence, unit_name  # noqa: E402

failures = []


def check(cond, msg):
    print(("  ok  " if cond else "  FAIL ") + msg)
    if not cond:
        failures.append(msg)


def doc_text(path):
    d = Document(path)
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts += [c.text for c in row.cells]
    for s in d.sections:
        parts += [p.text for p in s.footer.paragraphs]
    return "\n".join(parts)


def build(data):
    tmp = tempfile.mkdtemp()
    dp, out = os.path.join(tmp, "d.json"), os.path.join(tmp, "o.docx")
    json.dump(data, open(dp, "w", encoding="utf-8"), ensure_ascii=False)
    r = subprocess.run([sys.executable, os.path.join(SKILL, "scripts", "build_udostoverenie.py"),
                        dp, "-o", out], capture_output=True, text=True)
    return (doc_text(out) if r.returncode == 0 else None), r


def test_forces():
    print("forces_sentence")
    check(unit_name("РС ПБЗН - Плевен") == "РСПБЗН – Плевен", "РС ПБЗН → РСПБЗН")
    check(unit_name("У ПБЗН - Сторгозия") == "У ПБЗН – Сторгозия", "У ПБЗН stays")
    check(forces_sentence("служ 3; 1 ПА; 1 - РС ПБЗН - Плевен;") ==
          "Произшествието е ликвидирано от екип на РСПБЗН – Плевен с 1 бр. пожарен автомобил и 3 служители.",
          "single unit sentence")
    check(forces_sentence("служ 6; 2 ПА; 1 - РС ПБЗН - Плевен; 1 - У ПБЗН - Сторгозия;") ==
          "Произшествието е ликвидирано от екипи на РСПБЗН – Плевен и У ПБЗН – Сторгозия "
          "с 2 бр. пожарни автомобила и 6 служители.", "two units sentence")


def test_udostoverenie():
    print("udostoverenie")
    data = {
        "kind": "udostoverenie",
        "addressee": ["Г-Н ТЕСТ ТЕСТОВ", "ГР. ПЛЕВЕН", "ОБЩ. ПЛЕВЕН", "УЛ. „ТЕСТОВА“ №5"],
        "salutation": "УВАЖАЕМИ Г-Н ТЕСТОВ,",
        "body": ["Уведомявам Ви, че на 01.01.2026 г. в 10:00 ч. в ОЦ на РДПБЗН - Плевен е получен "
                 "сигнал за произшествие - пожар в лек автомобил.",
                 forces_sentence("служ 3; 1 ПА; 1 - РС ПБЗН - Плевен;"),
                 "Настоящето удостоверение се издава да послужи при необходимост."],
        "izgotvil": {"position": "инспектор IV ст.", "name": "Иван Дъбов", "date": "02.01.2026 г."},
        "copies": ["деловодство", "Тест Тестов"],
    }
    t, r = build(data)
    check(t is not None, f"builds ({r.stderr.strip()[-200:]})")
    if t is None:
        return
    check("УДОСТОВЕРЕНИЕ" in t and "ЗА ВЪЗНИКНАЛО ПРОИЗШЕСТВИЕ" in t, "title kept")
    check("Г-Н ТЕСТ ТЕСТОВ" in t and "УЛ. „ТЕСТОВА“ № 5" in t, "addressee lines, № spaced")
    check("УВАЖАЕМИ Г-Н ТЕСТОВ," in t, "salutation")
    check("РДПБЗН – Плевен е получен" in t and " - " not in t and "—" not in t, "en dash only")
    check("…………………..2026 г." in t, "year from izgotvil date")
    check("Стоян Кестенов" in t and "КОМИСАР" in t and "ДИРЕКТОР:" in t, "director from the staff list")
    check("Иван Дъбов" in t and "инспектор IV ст." in t and "02.01.2026 г." in t, "izgotvil")
    check("Отп: в 2 екз." in t and "Екз. № 2 – Тест Тестов" in t, "copies")
    for leftover in ["{{", "[Име", "[звание", "Ситроен", "ДРУЖБА"]:
        check(leftover not in t, f"no template leftovers: {leftover}")


def test_pridruzhitelno():
    print("pridruzhitelno")
    base = {
        "kind": "pridruzhitelno",
        "addressee": ["НАЧАЛНИКА НА", "РУ МВР – ЛЕВСКИ"],
        "attention": "Разследващ полицай Тест Тестов",
        "reference": "На Ваш рег. № 1р-1/01.01.2026 г.",
        "salutation": "УВАЖАЕМИ ГОСПОДИН НАЧАЛНИК,",
        "body": ["Във връзка с възникнал пожар на 01.01.2026 г. … приложено Ви изпращам заверени "
                 "копия на изготвените документи за възникналото произшествие."],
        "attachments": ["1. Телефонограма за произшествие № 1/01.01.2026 г. – 1 лист;",
                        "2. Статистически лист за произшествие № 1/01.01.2026 г. – 1 лист."],
        "izgotvil": {"date": "05.01.2026 г."},
        "copies": ["деловодство", "РУ МВР – Левски"],
    }
    t, r = build(base)
    check(t is not None, f"builds ({r.stderr.strip()[-200:]})")
    if t is None:
        return
    check("НАЧАЛНИКА НА" in t and "РУ МВР – ЛЕВСКИ" in t, "addressee rows")
    check("На вниманието на:" in t and "Разследващ полицай Тест Тестов" in t, "attention row")
    check("На Ваш рег. № 1р-1/01.01.2026 г." in t, "reference row")
    check(t.count("ПРИЛОЖЕНИЕ:") == 1 and "2. Статистически лист" in t, "attachments, label once")
    for leftover in ["{{", "Томов", "Комарево", "256р", "МИТРОПОЛИЯ"]:
        check(leftover not in t, f"no template leftovers: {leftover}")
    t2, _ = build(dict(base, attention="", reference=""))
    check(t2 is not None and "На вниманието на:" not in t2 and "На Ваш рег." not in t2,
          "empty attention/reference rows dropped")
    t3, r3 = build(dict(base, salutation=""))
    check(t3 is not None, "empty salutation still builds")


if __name__ == "__main__":
    test_forces()
    test_udostoverenie()
    test_pridruzhitelno()
    print("\nALL PASS" if not failures else f"\n{len(failures)} FAILED")
    sys.exit(1 if failures else 0)
