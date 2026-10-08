#!/usr/bin/env python3
"""Smoke test for build_spravka.py. Run from the skill folder: python3 tests/test_build.py"""
import json
import os
import subprocess
import sys
import tempfile

from docx import Document

SKILL = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

# the tests never touch the network: no key, and an invented staff list instead of the real one
os.environ["FIRE_SERVICE_NO_KEY"] = "1"
os.environ["FIRE_SERVICE_STAFF_CSV"] = os.path.abspath(os.path.join(SKILL, "..", "..", "..", "..", "tests", "fixtures", "staff.csv"))
failures = []


def check(cond, msg):
    print(("  ok  " if cond else "  FAIL ") + msg)
    if not cond:
        failures.append(msg)


def text(path):
    d = Document(path)
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts += [c.text for c in row.cells]
    return "\n".join(parts)


def body_alignments(data):
    """Alignment of every body paragraph (intro, headings, text) of a built справка."""
    tmp = tempfile.mkdtemp()
    dp, out = os.path.join(tmp, "d.json"), os.path.join(tmp, "s.docx")
    json.dump(data, open(dp, "w", encoding="utf-8"), ensure_ascii=False)
    subprocess.run([sys.executable, os.path.join(SKILL, "scripts", "build_spravka.py"), dp, "-o", out],
                   check=True, capture_output=True)
    ps = Document(out).paragraphs
    start = next(i for i, p in enumerate(ps) if p.text.startswith("Днес"))
    return {str(p.alignment) for p in ps[start:] if p.text.strip()}


def build(data):
    tmp = tempfile.mkdtemp()
    dp, out = os.path.join(tmp, "d.json"), os.path.join(tmp, "s.docx")
    json.dump(data, open(dp, "w", encoding="utf-8"), ensure_ascii=False)
    subprocess.run([sys.executable, os.path.join(SKILL, "scripts", "build_spravka.py"), dp, "-o", out],
                   check=True, capture_output=True)
    return text(out)


if __name__ == "__main__":
    base = {
        "intro": "Днес, 01.08.2025 г., в гр.Плевен, Община Плевен, подписаният Иван Дъбов – "
                 "инспектор ІV ст. в група „ОЦ” при РДПБЗН – Плевен, след ликвидиране на възникнал на "
                 "01.08.2025 г. пожар в лек автомобил, в гр. Плевен, и направения оглед установих следното:",
        "section1": ["Тестов абзац едно.", "Контейнерът е на ул. „Климент Охридски“ №4."],
        "section2": ["При пристигане на мястото на произшествието заварихме, че автомобилът гори."],
        "section3": ["Огнището на пожара е в двигателния отсек.",
                     "На основание гореизложеното считам, че най-вероятната непосредствена причина за "
                     "възникване на пожара е късо съединение."],
        "signer_title": "инспектор IV ст.",
        "signer_name": "Иван Дъбов",
    }
    print("build_spravka.py")
    t = build(dict(base, legal_basis="new"))
    for h in ["1. Обстановка, предшествуваща", "2. Обстановка по време", "3. Обстоятелства и факти"]:
        check(h in t, f"heading '{h}'")
    check("8121з – 1055/17.08.2021" in t, "legal basis new")
    check("НАЧАЛНИКА НА" in t and "РСПБЗН – ПЛЕВЕН" in t, "default addressee")
    check("ГЛАВЕН ИНСПЕКТОР" in t and "ВАСИЛ СМЪРЧЕВ" in t and t.index("ГЛАВЕН ИНСПЕКТОР") > t.index("РСПБЗН – ПЛЕВЕН"), "chief of РСПБЗН – Плевен in addressee")
    check("Иван Дъбов" in t, "signer")
    check("Охридски“ № 4." in t and "№4" not in t, "№ followed by a space")
    check("—" not in t and " - " not in t, "only en dash as a dash")
    check("Тестов абзац едно." in t and "късо съединение" in t, "body text")
    t_old = build(dict(base, legal_basis="old"))
    check("Iз – 2775/28.10.2011" in t_old, "legal basis old")
    check("вероятната причина за пожара:" in t and "причина за произшествието" not in t,
          "heading 3 wording follows чл. 20, ал. 3")
    check("РАЙОННА СЛУЖБА „ПОЖАРНА БЕЗОПАСНОСТ И ЗАЩИТА НА НАСЕЛЕНИЕТО” – ПЛЕВЕН" in t, "default unit header")
    t_unit = build(dict(base, unit="Левски"))
    check("РАЙОННА СЛУЖБА „ПОЖАРНА БЕЗОПАСНОСТ И ЗАЩИТА НА НАСЕЛЕНИЕТО” – ЛЕВСКИ" in t_unit, "unit header")
    check("РСПБЗН – ЛЕВСКИ" in t_unit and "РСПБЗН – ПЛЕВЕН" not in t_unit
          and "„ПБЗН” - ПЛЕВЕН" not in t_unit, "unit addressee, no Плевен left in header/addressee")
    check("ПЕТЪР БАДЕМОВ" in t_unit and "ВАСИЛ СМЪРЧЕВ" not in t_unit, "chief follows unit")
    check("ЧЕРВЕН БРЯГ" in build(dict(base, unit="ЧЕРВЕН БРЯГ")) and "НИКОЛАЙ ШИПКОВ" in build(dict(base, unit="ЧЕРВЕН БРЯГ")), "chief for two-word unit")
    check("Рег. № ...................., екз. № ......" in t and ".............2025 г." in t,
          "registration block, year from the intro date")
    check(".............2026 г." in build(dict(base, reg_year="2026")), "reg_year overrides the year")
    check("гр. Плевен" in t and "гр. Левски" in t_unit, "place follows unit")
    check("гр. Долна Митрополия" in build(dict(base, city="Долна Митрополия")), "city overrides the place")
    check("гр. Плевен\nИЗГОТВИЛ:\t\t\n\n(инспектор IV ст. Иван Дъбов)" in t,
          "closing block: place and ИЗГОТВИЛ on one line, tabs for the signature, signer below")
    base_def = {k: v for k, v in base.items() if not k.startswith("signer_")}
    check("([звание] [Име Фамилия])" in build(base_def), "no signer given → placeholders, never somebody's name")
    check("(инспектор VI ст. Красимира Лозева)" in
          build(dict(base, signer_title="ИНСПЕКТОР VI СТ.", signer_name="Красимира Лозева")),
          "upper-case title is lowered, Roman numeral kept")
    check("(мл. експерт Георги Малинов)" in build(dict(base, signer_title="мл. експерт", signer_name="Георги Малинов")),
          "mixed-case signer title is kept as given")
    t_cb = build(dict(base, unit="ЧЕРВЕН БРЯГ", signer_title="ВПД НАЧАЛНИК НА РСПБЗН – ЧЕРВЕН БРЯГ", signer_name="Иван Дъбов"))
    check("(ВПД началник на РСПБЗН – Червен бряг Иван Дъбов)" in t_cb,
          "abbreviations keep their capitals; the place is spelled officially – „Червен бряг“")
    check("гр. Червен бряг" in t_cb and "Червен Бряг" not in t_cb, "„гр. Червен бряг“, never „Червен Бряг“")
    check("(началник на РСПБЗН – Долна Митрополия Иван Дъбов)" in
          build(dict(base, signer_title="НАЧАЛНИК НА РСПБЗН – ДОЛНА МИТРОПОЛИЯ", signer_name="Иван Дъбов")),
          "a two-word place with both words capitalised stays so")
    check(body_alignments(base) == {"JUSTIFY (3)"}, "body text is justified")
    t_nc = build(dict(base, addressee_chief=False))
    check("ВАСИЛ СМЪРЧЕВ" not in t_nc and "РСПБЗН – ПЛЕВЕН" in t_nc, "addressee_chief false")
    print("\nALL PASS" if not failures else f"\n{len(failures)} FAILED")
    sys.exit(1 if failures else 0)
