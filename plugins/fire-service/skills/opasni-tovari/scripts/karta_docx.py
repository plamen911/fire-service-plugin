#!/usr/bin/env python3
"""
karta_docx.py — the dangerous goods card as a printable A4 .docx (and, with --pdf, a PDF next to it).

    karta_docx.py --un 1203 -o /mnt/user-data/outputs/ADR_1203_Benzin.docx
    karta_docx.py --tabela "265/1017" --vid tech --kolichestvo golyamo --vreme nosht \
                  --vyatar slab --sad avtocisterna -o ADR_1017_Hlor_proizshestvie.docx --pdf

Without --vid it is the reference card (the whole ERI card or ERG guide + distances); with --vid,
--kolichestvo and --vreme it is the incident card (zones first, then the actions that apply).
The content is exactly what adr.py prints – this script only lays it out.
"""
import argparse
import datetime
import os
import re
import shutil
import subprocess
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import adr  # noqa: E402

try:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
except ImportError:
    sys.exit("Липсва python-docx: pip install python-docx")

ASSETS = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "assets", "etiketi")
ORANGE = "F39200"
FONT = "Arial"


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def borders(cell, size=12, color="000000", sides=("top", "left", "bottom", "right")):
    tcPr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for side in sides:
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(size)); e.set(qn("w:color"), color)
        b.append(e)
    tcPr.append(b)


def no_table_borders(table):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), "nil")
        b.append(e)
    tblPr.append(b)


def run(p, text, size=9.5, bold=False, italic=False, color=None):
    r = p.add_run(text)
    r.font.name = FONT
    r._element.rPr.rFonts.set(qn("w:cs"), FONT); r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    r.font.size = Pt(size); r.bold = bold; r.italic = italic
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    return r


def para(container, space_after=1.5, align=None, left=0.0):
    p = container.add_paragraph()
    f = p.paragraph_format
    f.space_before = Pt(0); f.space_after = Pt(space_after); f.line_spacing = 1.0
    if left:
        f.left_indent = Cm(left)
    if align:
        p.alignment = align
    return p


def label_files(etiketi):
    out = []
    for group in etiketi.split(", ")[:1]:
        for code in group.split("+"):
            code = code.strip()
            key = "7" if code.startswith("7") and code != "7E" else ("9" if code == "9A" else code)
            f = os.path.join(ASSETS, key + ".png")
            if os.path.isfile(f) and f not in out:
                out.append(f)
    return out[:4]


def header_block(doc, r, title):
    t = doc.add_table(rows=1, cols=3)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    no_table_borders(t)
    widths = (Cm(4.2), Cm(9.6), Cm(4.2))
    for c, w in zip(t.rows[0].cells, widths):
        c.width = w
    left, mid, right = t.rows[0].cells
    # the orange plate
    plate = left.add_table(rows=2, cols=1)
    plate.autofit = False
    for i, text in enumerate(((r["hin"].split(",")[0].strip() if r["hin"] else ""), r["un"])):
        c = plate.rows[i].cells[0]
        c.width = Cm(3.8)
        shade(c, ORANGE); borders(c, 18)
        p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(1); p.paragraph_format.space_after = Pt(1)
        run(p, text or " ", size=24, bold=True)
    left.paragraphs[0].paragraph_format.space_after = Pt(0)
    # names
    p = mid.paragraphs[0]; p.paragraph_format.space_after = Pt(2)
    run(p, title, size=8.5, bold=True, color="7F7F7F")
    p = para(mid, 2); run(p, f"UN {r['un']} – {r['ime']}", size=13, bold=True)
    p = para(mid, 3); run(p, r["ime_en"], size=8.5, italic=True)
    p = para(mid, 0)
    bits = [f"Клас {adr.klas_text(r)}"]
    if r["kod"]:
        bits.append(f"кл. код {r['kod']}")
    if r["og"]:
        bits.append(f"оп. група {r['og']}")
    if r["etiketi"]:
        bits.append(f"етикети {r['etiketi']}")
    run(p, "; ".join(bits), size=8.5)
    # labels
    p = right.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    files = label_files(r["etiketi"])
    size = Cm(2.0) if len(files) <= 2 else Cm(1.55)
    for f in files:
        p.add_run().add_picture(f, width=size)
    if r["hin"]:
        p = para(doc, 3); p.paragraph_format.space_before = Pt(4)
        for h in [x.strip() for x in r["hin"].split(",")]:
            run(p, f"Номер за опасност {h}: ", size=9.5, bold=True); run(p, adr.hin_decode(h) + ". ", size=9.5)


def inline(p, text, size=9.5):
    """**bold** and `code` inside a line."""
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if part.startswith("**"):
            run(p, part[2:-2], size=size, bold=True)
        elif part.startswith("`"):
            run(p, part[1:-1], size=size - 1)
        elif part:
            run(p, part, size=size)


def body(doc, lines):
    zone = False
    for ln in lines:
        if not ln.strip():
            continue
        if ln.startswith("## ") or ln.startswith("_") and ln.endswith("_") and ln.count("_") == 2 and lines.index(ln) < 3:
            continue                                      # title and English name are in the header block
        if ln.startswith("- Оранжева табела") or ln.startswith("- Номер за опасност") or ln.startswith("- Клас ") \
                or ln.startswith("- Етикети") or ln.startswith("Табела ") or re.match(r"UN \d{4} \(без номер", ln):
            continue
        if ln.startswith("### "):
            title = ln[4:]
            zone = "Зони" in title or title.startswith("Разстояния")
            t = doc.add_table(rows=1, cols=1); t.autofit = False
            c = t.rows[0].cells[0]; c.width = Cm(18.0)
            shade(c, "C00000" if zone else "404040")
            p = c.paragraphs[0]; p.paragraph_format.space_before = Pt(1); p.paragraph_format.space_after = Pt(1)
            run(p, title.upper() if zone else title, size=10, bold=True, color="FFFFFF")
            continue
        if ln.startswith("**") and ln.endswith("**"):
            p = para(doc, 0.5); p.paragraph_format.space_before = Pt(2)
            run(p, ln[2:-2], size=9, bold=True)
        elif ln.startswith("  - "):
            p = para(doc, 1, left=0.9); inline(p, "– " + ln[4:], 9 if not zone else 10)
        elif ln.startswith("- "):
            p = para(doc, 0.5, left=0.35); p.paragraph_format.first_line_indent = Cm(-0.3)
            inline(p, "• " + ln[2:], 10 if zone else 9)
            if zone and not ln.startswith(("- „", "- За изолирането")) and re.search(r"изолирайте|защитете|евакуация", ln.lower()):
                for r_ in p.runs:
                    r_.bold = True
        elif ln.startswith("> "):
            p = para(doc, 2, left=0.35); run(p, ln[2:], size=9.5, bold=True)
        elif ln.startswith("_") and ln.endswith("_"):
            p = para(doc, 1); p.paragraph_format.space_before = Pt(2); run(p, ln.strip("_"), size=9.5, italic=True, bold=True)
        elif ln.startswith("Източници:") or ln.startswith("Справочна информация"):
            p = para(doc, 1); p.paragraph_format.space_before = Pt(4); run(p, ln, size=7.5, color="595959")
        else:
            ln = re.sub(r" Пълната карта \(.*$", "", ln)      # the pointer to the script means nothing on paper
            p = para(doc, 2); inline(p, ln, 9)


def build(r, lines, out, title):
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(1.5); s.top_margin = Cm(1.0); s.bottom_margin = Cm(1.0)
    st = doc.styles["Normal"]; st.font.name = FONT; st.font.size = Pt(9.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    for p0 in list(doc.paragraphs):
        p0._element.getparent().remove(p0._element)
    header_block(doc, r, title)
    body(doc, lines)
    p = para(doc, 0, align=WD_ALIGN_PARAGRAPH.RIGHT)
    run(p, "Изготвено на " + datetime.date.today().strftime("%d.%m.%Y") + " г. със скила „Опасни товари“ (fire-service).",
        size=7.5, color="595959")
    doc.save(out)


def to_pdf(docx_path):
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        return None
    outdir = os.path.dirname(os.path.abspath(docx_path))
    subprocess.run([exe, "--headless", "--convert-to", "pdf", "--outdir", outdir, docx_path],
                   capture_output=True, text=True, timeout=180)
    pdf = os.path.splitext(docx_path)[0] + ".pdf"
    return pdf if os.path.isfile(pdf) else None


def main(argv=None):
    p = argparse.ArgumentParser(description="Карта за опасен товар като .docx (и PDF)")
    adr.add_pick(p)
    p.add_argument("--vid", choices=list(adr.VID)); p.add_argument("--kolichestvo", choices=["malko", "golyamo"])
    p.add_argument("--vreme", choices=["den", "nosht"]); p.add_argument("--vyatar", choices=list(adr.VYATAR))
    p.add_argument("--sad", choices=list(adr.SAD)); p.add_argument("--voda", action="store_true")
    p.add_argument("-o", "--output", required=True)
    p.add_argument("--pdf", action="store_true", help="also write a PDF next to the .docx (needs LibreOffice)")
    a = p.parse_args(argv)
    r = adr.need_one(a)
    if a.vid:
        if not (a.kolichestvo and a.vreme):
            sys.exit("За карта за произшествие са нужни --vid, --kolichestvo и --vreme.")
        lines, _ = adr.incident_text(r, a)
        title = "ОПАСЕН ТОВАР – КАРТА ЗА ПРОИЗШЕСТВИЕ"
    else:
        lines = adr.card_text(r, a.erg or "")
        title = "ОПАСЕН ТОВАР – СПРАВОЧНА КАРТА"
    if a.note:
        lines = lines[:3] + [a.note, ""] + lines[3:]
    os.makedirs(os.path.dirname(os.path.abspath(a.output)) or ".", exist_ok=True)
    build(r, lines, a.output, title)
    print("DOCX:", a.output)
    if a.pdf:
        pdf = to_pdf(a.output)
        print("PDF:", pdf if pdf else "не е създаден – няма LibreOffice в средата; предайте .docx файла")


if __name__ == "__main__":
    main()
