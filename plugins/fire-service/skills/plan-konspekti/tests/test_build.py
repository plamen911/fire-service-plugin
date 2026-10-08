#!/usr/bin/env python3
"""Smoke tests for the scripts of plan-konspekti. Run from the skill folder:
    python3 tests/test_build.py
"""
import json
import os
import subprocess
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
SKILL = os.path.dirname(HERE)
os.environ["FIRE_SERVICE_NO_KEY"] = "1"   # the tests never touch the network
sys.path.insert(0, os.path.join(SKILL, "scripts"))

from docx import Document  # noqa: E402
from PIL import Image  # noqa: E402

failures = []
APPROVER = {"lines": ["ДИРЕКТОР НА", "РДПБЗН – ПЛЕВЕН:", "КОМИСАР"], "name": "Георги Дъбов"}
SIGN = {"date": "20.10.2026 г.", "name": "инспектор Иван Петров", "position": "Инспектор IV ст. в група ОЦ"}


def check(cond, msg):
    print(("  ok  " if cond else "  FAIL ") + msg)
    if not cond:
        failures.append(msg)


def doc_text(path):
    d = Document(path)
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts += [c.text for c in row.cells]
    return "\n".join(parts), d


def no_bold(path):
    """No bold run in the body, headers and footers (unused built-in styles do not count)."""
    import re
    import zipfile
    with zipfile.ZipFile(path) as z:
        parts = [n for n in z.namelist() if re.match(r"word/(document|header\d*|footer\d*)\.xml$", n)]
        return not any(re.search(r"<w:b(/| )", z.read(n).decode("utf-8", "ignore")) for n in parts)


def run(script, data, tmp, *extra):
    dp, out = os.path.join(tmp, "d.json"), os.path.join(tmp, "o.docx")
    with open(dp, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False)
    r = subprocess.run([sys.executable, os.path.join(SKILL, "scripts", script), dp, "-o", out,
                        "--uprazhneniya", os.path.join(HERE, "fixtures", "uprazhneniya.json"), *extra],
                       capture_output=True, text=True)
    if r.returncode == 0 and "DOC_CHECK" in r.stderr:   # the finished document must keep the common rules
        check(False, "DOC_CHECK: " + r.stderr.strip()[:200])
    return r, out


def test_oc():
    print("build_konspekt: oc")
    tmp = tempfile.mkdtemp()
    Image.new("RGB", (200, 100), "white").save(os.path.join(tmp, "f.png"))
    data = {"kind": "oc", "data": "27.10.2026 г.", "plan_grafik": "рег. № 947р-0000/15.12.2025 г.",
            "approver": APPROVER, "izgotvil": SIGN, "tema": "Гасене на пожари в житни масиви",
            "tsel": "Опресняване на знанията на служителите.", "zvena": ["РСПБЗН – Левски", "РСПБЗН – Белене"],
            "materialno": ["Правила №1", "Наредба №2"],
            "izlozhenie": ["## 1. Увод", "Текст - с **получер**.", "> Чл. 42. Цитат.", "- точка",
                           {"fig": "f.png", "caption": "Фиг. 1. Проба", "height_cm": 3}],
            "tema2": {"uprazhnenie": "1.1"}}
    r, out = run("build_konspekt.py", data, tmp)
    check(r.returncode == 0, "builds" + ("" if r.returncode == 0 else ": " + r.stderr[-300:]))
    if r.returncode:
        return
    info = json.loads(r.stdout)
    text, d = doc_text(out)
    check("П Л А Н – К О Н С П Е К Т" in text, "spaced title")
    check(text.index("МАТЕРИАЛНО ОСИГУРЯВАНЕ") < text.index("ХОД НА ЗАНЯТИЕТО:") < text.index("1. Увод"),
          "the heading that opens the exposition")
    check("РЕГИОНАЛНА ДИРЕКЦИЯ „ПОЖАРНА БЕЗОПАСНОСТ И ЗАЩИТА НА НАСЕЛЕНИЕТО” – ПЛЕВЕН" in text
          and "„ПБЗН" not in text.split("УТВЪРЖДАВАМ")[0], "letterhead: the directorate written in full")
    head = d.tables[0]._tbl
    sizes = [r.font.size.pt for p in d.tables[0].cell(0, 0).paragraphs for r in p.runs]
    check(sizes == [14, 14, 12] and head.xpath("./w:tblPr/w:tblBorders/w:bottom[@w:val='single']")
          and head.xpath("./w:tblPr/w:tblW/@w:w") == ["10188"], "letterhead: 14/14/12 pt over a rule, 10188 dxa wide")
    check("рег. № 947р-0000/15.12.2025 г." in text and "през 2026 г." in text, "plan-grafik and year in the subtitle")
    check("Сградите на РСПБЗН – Левски и РСПБЗН – Белене" in text, "place from the units")
    check("Вдигане в контролна тревога" in text and "№ 8121з-1702/09.12.2022 г." in text, "exercise 1.1 from the data file")
    r2 = subprocess.run([sys.executable, os.path.join(SKILL, "scripts", "build_konspekt.py"), os.path.join(tmp, "d.json"),
                         "-o", os.path.join(tmp, "x.docx")], capture_output=True, text=True)
    check(r2.returncode != 0 and "NO_KEY" in r2.stderr, "without a key and without a file the exercises are not read → NO_KEY")
    check("Правила № 1" in text and "Текст – с получер." in text, "typography: „№ 1“ and the dash")
    check("Георги Дъбов" in text and "инспектор Иван Петров" in text, "approver and author")
    check(len(d.inline_shapes) == 1 and info["figuri"] == 1 and "Фиг. 1. Проба" in text, "one figure with a caption")
    check(info["preduprezhdeniya"] == [], "no warnings")
    check(no_bold(out), "no bold anywhere in the document")
    data["tema2"] = {"uprazhnenie": "2.13 А", "opisanie": True}
    r, out = run("build_konspekt.py", data, tmp)
    text = doc_text(out)[0] if r.returncode == 0 else ""
    check("Действия с хидравличен спасителен комплект" in text and "Резултат: до 6 точки по картата" in text,
          "exercise description and result scale from the methodology")
    data["tema2"] = {"uprazhnenie": "9.99"}
    r, _ = run("build_konspekt.py", data, tmp)
    check(r.returncode != 0 and "BAD_INPUT" in r.stderr, "unknown exercise is refused")


def test_seminar():
    print("build_konspekt: seminar")
    tmp = tempfile.mkdtemp()
    data = {"kind": "seminar", "year": "2026", "approver": APPROVER, "izgotvil": SIGN,
            "tema": "Спасителни въжета", "tsel": "Опресняване на знанията.", "materialno": ["Правила"],
            "izlozhenie": ["Абзац."]}
    r, out = run("build_konspekt.py", data, tmp)
    check(r.returncode == 0, "builds")
    if r.returncode:
        return
    text, _ = doc_text(out)
    check("I. Тема: Спасителни въжета." in text and "VI. Материално осигуряване: Правила" in text, "fields I–VI")
    check("Рег. №" in text and "ТЕМА 2" not in text, "registration lines, no second topic")
    check("VII. Ход на занятието:" in text, "the exposition is the next numbered point")
    check("1 учебен час на" in text, "time without a date keeps a blank")
    data["izlozhenie"] = []
    r, _ = run("build_konspekt.py", data, tmp)
    check("няма изложение" in r.stdout, "warns when there is no exposition")


def test_psp():
    print("build_konspekt: psp")
    tmp = tempfile.mkdtemp()
    data = {"kind": "psp", "year": "2026", "zveno": "РСПБЗН – Левски", "uprazhnenie": "1.9 Б",
            "approver": {"lines": ["НАЧАЛНИК НА", "РСПБЗН – ЛЕВСКИ:", "ГЛАВЕН ИНСПЕКТОР"], "name": "Георги Дъбов"},
            "izgotvil": SIGN, "merki": ["Облекло и оборудване по Вариант № 2."], "minuti": {"osnovna": 30}}
    r, out = run("build_konspekt.py", data, tmp)
    check(r.returncode == 0, "builds" + ("" if r.returncode == 0 else ": " + r.stderr[-300:]))
    if r.returncode:
        return
    text, d = doc_text(out)
    check("за провеждане на занятие по ПСП със служителите от РСПБЗН – Левски" in text, "subtitle of Приложение № 4")
    check("ТЕМА: Упражнение № 1.9 Б" in text and "МЕРКИ ЗА БЕЗОПАСНОСТ И ЗДРАВЕ: Облекло" in text, "topic from the exercise, safety measures")
    check("Подготвителна част" in text and "до 10 минути" in text and "30 минути" in text and "до 5 минути" in text, "three parts with minutes")
    check("Резултат: 6 точки ≤ 45 секунди" in text, "main part: the exercise with its scale")
    check("ДАТА И ЧАСОВИ ИНТЕРВАЛ" in text and len(d.tables[-1].rows) == 7, "table of the sessions with 6 empty rows")
    check(no_bold(out), "no bold anywhere in the document")
    del data["uprazhnenie"]
    r, _ = run("build_konspekt.py", data, tmp)
    check(r.returncode != 0 and "BAD_INPUT" in r.stderr, "neither an exercise nor a topic → refused")


def test_otchet():
    print("build_otchet")
    tmp = tempfile.mkdtemp()
    data = {"data": "27.10.2026 г.", "tema1": "Гасене на пожари в житни масиви", "tema2": {"uprazhnenie": "1.1"},
            "izgotvil": SIGN,
            "zvena": [{"zveno": "РСПБЗН – Левски", "prisastvali": 3, "sekundi": 72, "tochki": 4},
                      {"zveno": "РСПБЗН – Белене", "prisastvali": 1, "rezultat": "Без грешки."}]}
    r, out = run("build_otchet.py", data, tmp)
    check(r.returncode == 0, "builds" + ("" if r.returncode == 0 else ": " + r.stderr[-300:]))
    if r.returncode == 0:
        text, _ = doc_text(out)
        check("О Т Ч Е Т" in text and "РСПБЗН – Левски и РСПБЗН – Белене" in text, "title and units")
        check("е 72 сек., което е за 4 точки." in text and "Без грешки." in text, "results")
        check("3 служители от състава" in text and "1 служител от състава" in text, "attendance")
        check("Екз. № 1 – деловодство" in text and "Инспектор IV ст. в група ОЦ" in text, "footer block")
        check(no_bold(out), "no bold anywhere in the document")
    del data["zvena"][0]["tochki"]
    r, out = run("build_otchet.py", data, tmp)
    check(r.returncode == 0 and "е 72 сек., което е за 4 точки." in doc_text(out)[0], "points from the norm of exercise 1.1")
    data["zvena"][0]["sekundi"] = 81
    r, out = run("build_otchet.py", data, tmp)
    check(r.returncode == 0 and "е 81 сек., което е за 0 точки." in doc_text(out)[0], "slower than the norm → 0 points")
    data["tema2"] = {"uprazhnenie": "2.8"}
    r, _ = run("build_otchet.py", data, tmp)
    check(r.returncode != 0 and "NO_NORM" in r.stderr, "seconds for an exercise without a time norm → NO_NORM")
    data["zvena"][0] = {"zveno": "РСПБЗН – Левски", "prisastvali": 3, "tochki": 7}
    r, _ = run("build_otchet.py", data, tmp)
    check(r.returncode != 0 and "BAD_INPUT" in r.stderr, "more points than the card allows are refused")
    data["tema2"] = {"uprazhnenie": "1.1"}
    data["zvena"] = [{"zveno": "РСПБЗН – Плевен"}, {"zveno": "РСПБЗН – Кнежа"}]
    r, out = run("build_otchet.py", data, tmp)
    text = doc_text(out)[0] if r.returncode == 0 else ""
    import re
    got = [(int(a), int(b)) for a, b in re.findall(r"е (\d+) сек\., което е за (\d) точки\.", text)]
    check("ПРИСЪСТВАЛИ СЛУЖИТЕЛИ: 6 служители" in text and "ПРИСЪСТВАЛИ СЛУЖИТЕЛИ: 3 служители" in text,
          "attendance filled in by the usual size of the shift (Плевен 6, others 3)")
    check(len(got) == 2 and all(p in (5, 6) and s <= 70 for s, p in got) and "………" not in text
          and "по обичайното" in r.stdout, "result filled in: excellent or good, and reported")
    r2, out2 = run("build_otchet.py", data, tmp)
    check(doc_text(out2)[0] == text, "the filled-in values are the same on every build")
    data["popalni"] = False
    r, out = run("build_otchet.py", data, tmp)
    check(r.returncode == 0 and "ПРИСЪСТВАЛИ СЛУЖИТЕЛИ: ……… служители" in doc_text(out)[0]
          and "е ……… сек., което е за ……… точки." in doc_text(out)[0], "popalni: false leaves dots to fill in by hand")
    from build_otchet import points
    norm = [{"do_sek": 60, "tochki": 6}, {"do_sek": 75, "tochki": 4}]
    check(points(norm, 60) == 6 and points(norm, 72) == 4 and points(norm, 90) == 0, "points by a norm table")


def test_grafik():
    print("plan_grafik")
    fx = os.path.join(HERE, "fixtures", "plan_grafik_2026.md")
    cmd = [sys.executable, os.path.join(SKILL, "scripts", "plan_grafik.py")]
    r = subprocess.run(cmd + ["mesec", "2026-10", "--ime", "Иван Петров", "--file", fx], capture_output=True, text=True)
    check(r.returncode == 0, "reads a saved file")
    if r.returncode == 0:
        d = json.loads(r.stdout)
        z = d["zanyatiya"]
        check(len(z) == 1 and z[0]["zvena"] == ["РСПБЗН – Левски", "РСПБЗН – Белене"], "one lesson of this person, two units")
        check(z[0]["data"] == "27.10.2026 г." and d["plan_grafik"] == "рег. № 947р-0000/15.12.2025 г.", "date and registration")
    r = subprocess.run(cmd + ["godina", "2026", "--file", fx], capture_output=True, text=True)
    check(r.returncode == 0 and len(json.loads(r.stdout)["zanyatiya"]) == 3, "the whole year")
    r = subprocess.run(cmd + ["mesec", "2026-10"], capture_output=True, text=True)
    check(r.returncode != 0 and "NO_KEY" in r.stderr, "without a key → NO_KEY")


def test_figuri():
    print("figuri")
    try:
        from pptx import Presentation
        from pptx.util import Cm
    except ImportError:
        print("  skip (python-pptx is not installed)")
        return
    tmp = tempfile.mkdtemp()
    for name, size in (("a.png", (60, 120)), ("b.png", (80, 120))):
        Image.new("RGB", size, "black").save(os.path.join(tmp, name))
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Булин"
    s.shapes.add_picture(os.path.join(tmp, "b.png"), Cm(10), Cm(5))
    s.shapes.add_picture(os.path.join(tmp, "a.png"), Cm(2), Cm(5))
    prs.save(os.path.join(tmp, "p.pptx"))
    r = subprocess.run([sys.executable, os.path.join(SKILL, "scripts", "figuri.py"), os.path.join(tmp, "p.pptx"),
                        "-o", os.path.join(tmp, "fig")], capture_output=True, text=True)
    check(r.returncode == 0, "reads a presentation")
    if r.returncode == 0:
        d = json.loads(r.stdout)
        sl = d["slides"][0]
        check(d["figuri"] == 1 and sl["izobrazheniya"] == 2 and "Булин" in sl["tekst"], "one row of two pictures, with the text")
        check(sl["shirina_px"] == (60 + 80) * 2 + 24 and os.path.isfile(sl["fig"]), "pictures side by side, enlarged")


if __name__ == "__main__":
    for t in (test_oc, test_seminar, test_psp, test_otchet, test_grafik, test_figuri):
        t()
    print(f"\n{len(failures)} failures" if failures else "\nall ok")
    sys.exit(1 if failures else 0)
