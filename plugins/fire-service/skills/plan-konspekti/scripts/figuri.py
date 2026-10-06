#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
figuri.py — изважда текста и изображенията от презентация (.pptx), за да влязат като фигури в
план-конспекта. За всеки слайд с изображения прави ЕДИН ред – изображенията му, наредени отляво
надясно върху бял фон, уголемени за печат (много широките – едно под друго) – и го записва като PNG.

    python3 scripts/figuri.py prezentatsiya.pptx -o fig

Отпечатва JSON: {"slaydove": N, "figuri": M, "slides": [{"n": 2, "tekst": ["…"], "belezhki": "…",
"izobrazheniya": 2, "fig": "fig/s02.png", "shirina_px": 860, "visochina_px": 908}, …]}.
Полето "fig" отива направо в "izlozhenie" на data.json за build_konspekt.py:
    {"fig": "fig/s02.png", "caption": "Фиг. 1. …", "height_cm": 4.5}

Грешки (на stderr, код ≠ 0):
    BAD_INPUT: ...   — файлът не е .pptx или не се чете
"""
import argparse
import io
import json
import os
import sys

SCALE = 2          # the pictures in presentations are small; enlarge them for print
GAP = 24           # pixels of white between the pictures of one slide (after scaling)
WIDE = 5.0         # a row wider than this many heights is stacked vertically instead


def pictures(shapes):
    """Pictures of a slide, including those inside groups."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    out = []
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            out += pictures(sh.shapes)
        elif sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            out.append(sh)
    return out


def texts(shapes):
    out = []
    for sh in shapes:
        if getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip():
            out.append(sh.text_frame.text.strip())
        if getattr(sh, "has_table", False) and sh.has_table:
            for row in sh.table.rows:
                out.append(" | ".join(c.text.strip() for c in row.cells))
    return out


def row_image(pics):
    from PIL import Image
    parts = []
    for sh in sorted(pics, key=lambda s: (s.left or 0, s.top or 0)):
        im = Image.open(io.BytesIO(sh.image.blob)).convert("RGBA")
        bg = Image.new("RGBA", im.size, "white")
        bg.alpha_composite(im)
        w, h = bg.size
        box = (int(w * sh.crop_left), int(h * sh.crop_top), int(w * (1 - sh.crop_right)), int(h * (1 - sh.crop_bottom)))
        bg = bg.crop(box).convert("RGB")
        parts.append(bg.resize((bg.width * SCALE, bg.height * SCALE), Image.LANCZOS))
    wide = sum(p.width for p in parts) / max(p.height for p in parts) > WIDE
    if wide and len(parts) > 1:                      # long, low pictures: one under the other
        width, height = max(p.width for p in parts), sum(p.height for p in parts) + GAP * (len(parts) - 1)
        row, y = Image.new("RGB", (width, height), "white"), 0
        for p in parts:
            row.paste(p, ((width - p.width) // 2, y))
            y += p.height + GAP
        return row
    height = max(p.height for p in parts)
    width = sum(p.width for p in parts) + GAP * (len(parts) - 1)
    row = Image.new("RGB", (width, height), "white")
    x = 0
    for p in parts:
        row.paste(p, (x, (height - p.height) // 2))
        x += p.width + GAP
    return row


def main():
    ap = argparse.ArgumentParser(description="Текст и фигури от презентация (.pptx)")
    ap.add_argument("pptx")
    ap.add_argument("-o", "--output", default="fig", help="папка за фигурите (по подразбиране fig)")
    a = ap.parse_args()
    try:
        from pptx import Presentation
        prs = Presentation(a.pptx)
    except Exception as e:  # noqa: BLE001
        sys.exit(f"BAD_INPUT: презентацията не се чете ({type(e).__name__}: {e})"[:300])
    os.makedirs(a.output, exist_ok=True)
    slides, count = [], 0
    for n, slide in enumerate(prs.slides, 1):
        entry = {"n": n, "tekst": texts(slide.shapes), "belezhki": "", "izobrazheniya": 0, "fig": None}
        if slide.has_notes_slide:
            entry["belezhki"] = slide.notes_slide.notes_text_frame.text.strip()
        pics = pictures(slide.shapes)
        if pics:
            row = row_image(pics)
            path = os.path.join(a.output, f"s{n:02d}.png")
            row.save(path)
            count += 1
            entry.update(izobrazheniya=len(pics), fig=path, shirina_px=row.width, visochina_px=row.height)
        slides.append(entry)
    print(json.dumps({"slaydove": len(slides), "figuri": count, "slides": slides}, ensure_ascii=False, indent=1))


if __name__ == "__main__":
    main()
