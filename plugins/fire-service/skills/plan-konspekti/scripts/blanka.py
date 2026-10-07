#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
blanka.py — common building blocks of the documents of the skill plan-konspekti: the page, the
letterhead of РДПБЗН – Плевен, the „УТВЪРЖДАВАМ“ block, body paragraphs and the „ИЗГОТВИЛ“ block.
Imported by build_konspekt.py and build_otchet.py; not run on its own.

The look follows the approved documents of the directorate: A4, Times New Roman 12 pt, exact
line spacing, first-line indent 1.5 cm, the letterhead and the approval block in borderless tables.
NO BOLD anywhere (a rule of the directorate, to save toner – _shared/conventions.md): headings and
labels are plain text, emphasis is italic.
Typography follows _shared/conventions.md: „№ 4“ and the dash „–“.
"""
import os
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
FONT = "Times New Roman"
LINE = Pt(16)            # exact line spacing of the body
INDENT = Cm(1.5)         # first-line indent
# The letterhead – one look for every document of the plugin (_shared/conventions.md, „Шапка“):
# (text, size in pt, character spacing in twentieths of a point, space before in pt).
HEADER = [("МИНИСТЕРСТВО НА ВЪТРЕШНИТЕ РАБОТИ", 14, 32, 0),
          ("ГЛАВНА ДИРЕКЦИЯ „ПОЖАРНА БЕЗОПАСНОСТ И ЗАЩИТА НА НАСЕЛЕНИЕТО”", 14, -16, 6),
          ("РЕГИОНАЛНА ДИРЕКЦИЯ „ПОЖАРНА БЕЗОПАСНОСТ И ЗАЩИТА НА НАСЕЛЕНИЕТО” – ПЛЕВЕН", 12, -16, 6)]
HEADER_WIDTH = 10188     # dxa – wider than the text block, centered on it, as on the directorate's blank
PLACEHOLDER_NAME = "[Име Фамилия]"
PLACEHOLDER_RANK = "[ЗВАНИЕ]"

NO_RE = re.compile(r"№\s*(?=\S)")
DASH_RE = re.compile(r"(^|\s)-(?=\s)")
JUSTIFY, CENTER, LEFT = WD_ALIGN_PARAGRAPH.JUSTIFY, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT


def typo(text):
    """„№ 4“; the dash is „–“ (never „-“ between spaces, never „—“); closing quote is „“."""
    text = NO_RE.sub("№ ", str(text)).replace("—", "–").replace("”", "“")
    return DASH_RE.sub(lambda m: m.group(1) + "–", text)


def new_document(right_cm=2.0):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin, sec.right_margin = Cm(2.5), Cm(right_cm)
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(1.5)
    st = doc.styles["Normal"]
    st.font.name, st.font.size = FONT, Pt(12)
    rpr = st.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(a), FONT)
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "bg-BG")
    rpr.append(lang)
    st.paragraph_format.space_before = st.paragraph_format.space_after = Pt(0)
    return doc


def fmt(p, align=None, indent=False, line=LINE, keep=False):
    pf = p.paragraph_format
    pf.space_before = pf.space_after = Pt(0)
    if line:
        pf.line_spacing_rule, pf.line_spacing = WD_LINE_SPACING.EXACTLY, line
    if indent:
        pf.first_line_indent = INDENT
    if align is not None:
        p.alignment = align
    if keep:
        pf.keep_with_next = True
    return p


def run(p, text, italic=False, size=None):
    r = p.add_run(typo(text))
    r.italic = italic or None
    r.font.name = FONT
    if size:
        r.font.size = Pt(size)
    return r


def para(doc, text="", align=JUSTIFY, indent=True, italic=False, size=None, line=LINE, keep=False):
    """One paragraph. „*…*“ inside the text marks italic parts (the name of a knot, a term);
    „**…**“ is taken as the same – there is no bold in the documents."""
    p = fmt(doc.add_paragraph(), align, indent, line, keep)
    for i, part in enumerate(re.split(r"\*{1,2}(.+?)\*{1,2}", text)):
        if part:
            run(p, part, italic=italic or i % 2 == 1, size=size)
    return p


def blank(doc, n=1, line=LINE):
    for _ in range(n):
        fmt(doc.add_paragraph(), line=line)


def _borderless(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tbl_pr.append(borders)


def _cell(cell, lines, align=LEFT):
    """lines: [text | (text, {"size", "align"})]; the first one reuses the cell's paragraph."""
    for i, item in enumerate(lines):
        text, opt = (item, {}) if isinstance(item, str) else item
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        fmt(p, opt.get("align", align), line=None)
        if text:
            run(p, text, size=opt.get("size"))


def letterhead(doc, extra=None):
    """The letterhead: three centered lines over a thin rule, the same in every document of the
    plugin; extra = one more line under them (the unit, e.g. „РСПБЗН – Левски“)."""
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    tbl_pr = t._tbl.tblPr
    width = tbl_pr.find(qn("w:tblW"))
    if width is None:
        width = OxmlElement("w:tblW")
        tbl_pr.append(width)
    width.set(qn("w:w"), str(HEADER_WIDTH))
    width.set(qn("w:type"), "dxa")
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single" if edge == "bottom" else "nil")
        if edge == "bottom":
            el.set(qn("w:sz"), "4")
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), "auto")
        borders.append(el)
    after = [tbl_pr.find(qn(f"w:{n}")) for n in ("shd", "tblLayout", "tblCellMar", "tblLook")]
    after = [el for el in after if el is not None]
    if after:
        after[0].addprevious(borders)                             # schema order inside tblPr
    else:
        tbl_pr.append(borders)
    for col in t._tbl.tblGrid.findall(qn("w:gridCol")):
        col.set(qn("w:w"), str(HEADER_WIDTH))
    cell = t.cell(0, 0)
    cell.width = HEADER_WIDTH * 635
    lines = HEADER + ([(extra.upper(), 12, -16, 6)] if extra else [])
    for i, (text, size, spacing, before) in enumerate(lines):
        p = fmt(cell.paragraphs[0] if i == 0 else cell.add_paragraph(), CENTER, line=LINE)
        p.paragraph_format.space_before = Pt(before)
        r = p.add_run(text if i < len(HEADER) else typo(text))    # the blank's own quotes stay as they are
        r.font.name, r.font.size = FONT, Pt(size)
        sp = OxmlElement("w:spacing")
        sp.set(qn("w:val"), str(spacing))
        rpr = r._r.get_or_add_rPr()
        size_el = rpr.find(qn("w:sz"))
        if size_el is not None:
            size_el.addprevious(sp)                               # schema order: spacing before sz
        else:
            rpr.append(sp)
    blank(doc)


def grid(doc, header, rows, widths_cm, align=None):
    """A bordered table: header = column titles, rows = lists of cell texts (a cell may be a list of
    paragraphs). No bold; the header is centered."""
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    t._tbl.tblPr.append(layout)
    for col, w in zip(t.columns, widths_cm):
        col.width = Cm(w)
    for r, cells in enumerate([header] + rows):
        for c, value in enumerate(cells):
            cell = t.cell(r, c)
            cell.width = Cm(widths_cm[c])
            lines = value if isinstance(value, list) else [value]
            _cell(cell, [str(x) for x in lines] or [""], align=CENTER if r == 0 else (align or JUSTIFY))
    return t


def unit_chief_approver(unit):
    """The approval lines for the chief of a РСПБЗН/УПБЗН from the staff list, or None."""
    mod = staff_module()
    if mod is None:
        return None
    try:
        ch = mod.nachalnik(re.sub(r"^.*?[–-]\s*", "", unit))
    except BaseException:  # noqa: BLE001
        return None
    if not ch or not ch.get("ime"):
        return None
    post = ("ВПД " if ch.get("vpd") else "") + "НАЧАЛНИК НА"
    return {"lines": [post, (ch.get("zveno") or unit).upper() + ":", (ch.get("zvanie") or "").upper()], "name": ch["ime"]}


def approval(doc, approver, year, reg=False, width_cm=16.5):
    """The „УТВЪРЖДАВАМ“ block on the right; with reg=True the „Рег. №“ lines on the left."""
    left = ["Рег. № ........................, екз. № .......", f"..........................................{year} г."] if reg else []
    right = ["УТВЪРЖДАВАМ:"] + list(approver.get("lines") or []) + \
            [("", {}), (approver.get("name") or PLACEHOLDER_NAME, {"align": CENTER}), f"........................{year} г."]
    if approver.get("zamestvane"):
        right.append(approver["zamestvane"])
    t = doc.add_table(rows=1, cols=2)
    _borderless(t)
    t.autofit = False
    for cell, w in zip(t.rows[0].cells, (Cm(width_cm * 0.56), Cm(width_cm * 0.44))):
        cell.width = w
    _cell(t.cell(0, 0), left or [""])
    _cell(t.cell(0, 1), right)
    blank(doc, 2)


def title(doc, text, subtitle_lines):
    spaced = " ".join(text) if " " not in text.strip() else text   # „П Л А Н – К О Н С П Е К Т“
    p = fmt(doc.add_paragraph(), CENTER, line=None)
    run(p, spaced, size=16)
    blank(doc)
    for line in subtitle_lines:
        para(doc, line, align=CENTER, indent=False)
    blank(doc)


def signature(doc, date, lines, name, width_cm=16.5, left_extra=None):
    """„ИЗГОТВИЛ“ on the right; the date and „Отп. в 1 екз.“ on the left."""
    left = ([date, "Отп. в 1 екз."] if date else [""]) + list(left_extra or [])
    right = ["ИЗГОТВИЛ:"] + list(lines or []) + ["", (name or PLACEHOLDER_NAME, {"align": CENTER})]
    t = doc.add_table(rows=1, cols=2)
    _borderless(t)
    t.autofit = False
    for cell, w in zip(t.rows[0].cells, (Cm(width_cm * 0.5), Cm(width_cm * 0.5))):
        cell.width = w
    _cell(t.cell(0, 0), left)
    _cell(t.cell(0, 1), right)


def page_numbers(doc):
    """Centered page number in the footer (documents with an exposition run over several pages)."""
    p = fmt(doc.sections[0].footer.paragraphs[0], CENTER, line=None)
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        r = p.add_run()
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            el.text = text
        r._r.append(el)
        r.font.size = Pt(10)


def join_units(units):
    """["РСПБЗН – Левски", "РСПБЗН – Белене"] → „РСПБЗН – Левски и РСПБЗН – Белене“."""
    units = [typo(u).strip() for u in units if str(u).strip()]
    return units[0] if len(units) == 1 else ", ".join(units[:-1]) + " и " + units[-1] if units else ""


def staff_module():
    import importlib.util
    path = os.path.join(HERE, "..", "_shared", "scripts", "sluzhiteli.py")
    if not os.path.exists(path):
        return None
    try:
        spec = importlib.util.spec_from_file_location("sluzhiteli", path)
        mod = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(mod)
        return mod
    except Exception:  # noqa: BLE001  (no key, no network: the caller falls back to placeholders)
        return None


def director_approver():
    """The approval lines for the director of РДПБЗН – Плевен from the staff list, or placeholders."""
    mod = staff_module()
    d = None
    if mod is not None:
        try:
            d = mod.direktor()
        except BaseException:  # noqa: BLE001
            d = None
    if not d:
        return {"lines": ["ДИРЕКТОР НА", "РДПБЗН – ПЛЕВЕН:", PLACEHOLDER_RANK], "name": PLACEHOLDER_NAME}
    post = d["label"].rstrip(":")
    return {"lines": [f"{post} НА", "РДПБЗН – ПЛЕВЕН:", d["rank"]], "name": d["name"]}
