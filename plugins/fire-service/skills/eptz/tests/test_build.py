#!/usr/bin/env python3
"""Smoke tests for build_eptz_docx.py and build_smetka.py. Run from the skill folder:
    python3 tests/test_build.py
"""
import json
import os
import re
import subprocess
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
SKILL = os.path.dirname(HERE)

# the tests never touch the network: no key, and an invented staff list instead of the real one
os.environ["FIRE_SERVICE_NO_KEY"] = "1"
os.environ["FIRE_SERVICE_STAFF_CSV"] = os.path.abspath(os.path.join(SKILL, "..", "..", "..", "..", "tests", "fixtures", "staff.csv"))
sys.path.insert(0, os.path.join(SKILL, "scripts"))

from docx import Document  # noqa: E402
import openpyxl  # noqa: E402
from build_smetka import amount_words, apply_defaults, compute, words  # noqa: E402

FIXTURE = os.path.join(HERE, "fixtures", "sample_report.md")  # synthetic, not a real case
failures = []


def check(cond, msg):
    print(("  ok  " if cond else "  FAIL ") + msg)
    if not cond:
        failures.append(msg)


def doc_text(path):
    d = Document(path)
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts += [c.text for c in row.cells]
    for s in d.sections:
        parts += [p.text for p in s.footer.paragraphs]
    return "\n".join(parts)


def build(content, meta, out, final=False):
    cmd = [sys.executable, os.path.join(SKILL, "scripts", "build_eptz_docx.py"),
           "--content", content, "--meta", meta, "--out", out]
    if final:
        cmd.append("--final")
    subprocess.run(cmd, check=True, capture_output=True)


def test_eptz():
    print("build_eptz_docx.py")
    tmp = tempfile.mkdtemp()
    content = FIXTURE
    meta = {
        "zm_no": "123/2025", "zm_opis": "Първо РУ – Плевен", "dp_no": "1234/2025",
        "rp_grad": "Плевен", "prokuratura": "ОП – Плевен",
        "postanovlenie_no": "1234р-56789", "postanovlenie_date": "03.09.2026 г.",
        "naznachil": "разследващ полицай Тест Тестов", "expert_name": "Иван Дъбов", "expert_rank": "инспектор", "izgotvil_position": "инспектор IV ст.",
        "expert_position": "инспектор в група „ОЦ“ на сектор „ПГ и СД“ при РДПБЗН – Плевен",
        "izgotvil_date": "21.09.2026 г.", "copies": ["деловодство", "Първо РУ – Плевен"],
    }
    mp = os.path.join(tmp, "meta.json")
    json.dump(meta, open(mp, "w", encoding="utf-8"), ensure_ascii=False)

    draft, final = os.path.join(tmp, "draft.docx"), os.path.join(tmp, "final.docx")
    build(content, mp, draft)
    build(content, mp, final, final=True)
    t_d, t_f = doc_text(draft), doc_text(final)

    for h in ["I. Основание", "II. Представен", "III. Обстоятелства", "IV. Поставени", "V. Пожаро"]:
        check(h in t_f, f"section '{h}' present")
    check("Отговор на въпрос № 3" in t_f, "all 3 answers present")
    check("ЧЕРНОВА" not in t_d, "no ЧЕРНОВА notice without --final")
    check("ЧЕРНОВА" not in t_f, "no ЧЕРНОВА notice with --final")
    check(t_d == t_f, "--final changes nothing")
    check("ОП – Плевен" in t_f, "prokuratura override used")
    check("по описа на РП -" not in t_f, "no hardcoded РП when prokuratura is set")
    check("**" not in t_f, "no markdown bold left")
    check("[Източник" not in t_f, "citation markers stripped")
    check("1234/2025" in t_f and "123/2025" in t_f, "ДП/ЗМ numbers filled")
    check("Екз. № 2 – Първо РУ – Плевен" in t_f, "copies filled")

    # without prokuratura → РП – Плевен
    meta.pop("prokuratura")
    json.dump(meta, open(mp, "w", encoding="utf-8"), ensure_ascii=False)
    out = os.path.join(tmp, "rp.docx")
    build(content, mp, out, final=True)
    check("по описа на РП – Плевен" in doc_text(out), "default РП – {rp_grad}")
    nc = os.path.join(tmp, "no.md")
    open(nc, "w", encoding="utf-8").write(open(FIXTURE, encoding="utf-8").read().replace(
        "## I. Обстоятелства по делото.\n", "## I. Обстоятелства по делото.\n\nКонтейнерът е на ул. „Климент Охридски“ №4.\n", 1))
    build(nc, mp, out)
    t_no = doc_text(out)
    check("Охридски“ № 4." in t_no and "№4" not in t_no, "№ followed by a space")
    check("—" not in t_no and " - " not in t_no, "only en dash as a dash")


def test_smetka():
    print("build_smetka.py")
    check(words(21) == "двадесет и едно", "21 in words")
    check(words(125) == "сто двадесет и пет", "125 in words")
    check(words(1005) == "хиляда и пет", "1005 in words")
    check(words(2021) == "две хиляди двадесет и едно", "2021 in words")
    check(amount_words(97.61) == "Деветдесет и седем евро и шестдесет и един евроцента.", "97.61 EUR in words")
    check(amount_words(1.01) == "Едно евро и един евроцент.", "1.01 EUR in words")
    d = {"hours": 4, "a4_bw": 6, "ogled": False}
    check(compute(d) == 97.61, f"total for 4 h + 6 sheets = 97.61 (got {compute(d)})")
    dd = apply_defaults({"pages": 6, "copies": 2})
    check(dd["hours"] == 4 and dd["a4_bw"] == 12, "pages 6, copies 2 → 4 h, 12 sheets")
    check(apply_defaults({"pages": 6, "hours": 5})["hours"] == 5, "explicit hours win")
    tmp = tempfile.mkdtemp()
    data = dict(d, naznachil="разследващ полицай Тест", pri="Първо РУ – Плевен",
                delo_no="1234/2025", po_opisa_na="РП – Плевен", izvarshena_ot="Иван Дъбов", expert_title="инспектор", izgotvil_short="И. Дъбов",
                expert_position="инспектор IV ст. в група ОЦ към РДПБЗН – Плевен при ГДПБЗН – МВР",
                date="21.09.2026")
    dp = os.path.join(tmp, "s.json")
    json.dump(data, open(dp, "w", encoding="utf-8"), ensure_ascii=False)
    out = os.path.join(tmp, "s.xlsx")
    subprocess.run([sys.executable, os.path.join(SKILL, "scripts", "build_smetka.py"), dp, "-o", out],
                   check=True, capture_output=True)
    ws = openpyxl.load_workbook(out).active
    check("1234/2025" in ws["A7"].value and "Първо РУ – Плевен" in ws["A7"].value, "A7 filled")
    check(ws["D11"].value == 4 and ws["D16"].value == 6 and ws["D19"].value == 0, "quantities filled")
    check("Деветдесет и седем евро" in ws["A30"].value, "total in words written")
    yellow = [c.coordinate for r in ws.iter_rows() for c in r
              if c.fill is not None and c.fill.fill_type and str(c.fill.fgColor.rgb).upper() == "FFFFFF00"]
    check(not yellow, f"yellow fill cleared {yellow}")
    check("извършена от инспектор Иван Дъбов – инспектор IV ст. в група ОЦ" in ws["A7"].value,
          "expert title and position from the profile")
    check(ws["A43"].value == "/ инспектор И. Дъбов /", "„Изготвил“")
    check(ws["D44"].value == "КОМИСАР" and ws["D45"].value.strip() == "Стоян Кестенов", "director from the staff list")

    other = dict(data, izvarshena_ot="Георги Малинов", expert_title="старши пожарникар",
                 expert_position="старши пожарникар в РСПБЗН – Левски", izgotvil_short="Г. Малинов")
    json.dump(other, open(dp, "w", encoding="utf-8"), ensure_ascii=False)
    subprocess.run([sys.executable, os.path.join(SKILL, "scripts", "build_smetka.py"), dp, "-o", out],
                   check=True, capture_output=True)
    ws = openpyxl.load_workbook(out).active
    check("извършена от старши пожарникар Георги Малинов – старши пожарникар в РСПБЗН – Левски"
          in ws["A7"].value, "another expert's title and position in A7")
    check("РДПБЗН – Плевен" not in ws["A7"].value and "инспектор IV" not in ws["A7"].value,
          "no default position left for another expert")
    check(ws["A43"].value == "/ старши пожарникар Г. Малинов /", "another expert under „Изготвил“")


if __name__ == "__main__":
    test_eptz()
    test_smetka()
    print("\nALL PASS" if not failures else f"\n{len(failures)} FAILED")
    sys.exit(1 if failures else 0)
