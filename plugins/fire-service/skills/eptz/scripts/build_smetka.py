#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build_smetka.py — попълва СМЕТКА по чл. 23, ал. 3 от Наредба № Н-1/14.02.2023 г.
(възнаграждение за ЕПТЗ) върху шаблона assets/smetka_template.xlsx и я записва като .xlsx (по подразбиране).

Употреба:
    python3 build_smetka.py smetka.json -o /mnt/user-data/outputs/Smetka_....xlsx

smetka.json:
{
  "ekspertiza_no": "",                     // рег. № на експертизата — празно = остават точки
  "ekspertiza_year": "2026",
  "naznachil": "разследващ полицай Георги Георгиев",
  "pri": "Първо РУ – Плевен",
  "delo_vid": "ДП",                        // ДП / НОХД / гр.дело / адм.дело / изп.дело
  "delo_no": "1234/2025",
  "po_opisa_na": "РП – Плевен",
  "izvarshena_ot": "Иван Петров",            // име и фамилия на експерта (без бащино)
  "expert_title": "инспектор",               // пред името: „извършена от инспектор …“, „/ инспектор И. Петров /“
  "expert_position": "инспектор IV ст. в група ОЦ на сектор ПГ и СД към РДПБЗН – Плевен при ГДПБЗН – МВР",
                                           // длъжността след тирето – expert_position от профила на експерта
  "pages": 6,                              // страници на ЕПТЗ → часове труд = pages - 2 (ако няма "hours")
  "copies": 2,                             // екземпляри → листа А4 = pages × copies (ако няма "a4_bw")
  "hours": 4,                              // часове труд — по подразбиране pages - 2
  "a4_bw": 12,                             // отпечатани листа А4 черно-бяло — по подразбиране pages × copies
  "a4_color": 0,
  "a3_bw": 0,
  "ogled": false,                          // true → ред 3.4 (ръкавици, дезинфектант, маска) = 1
  "hours_weekend": 0,
  "hours_holiday": 0,
  "complex_addon": 0,                      // сума за особено сложна експертиза (EUR), по избор
  "izgotvil_short": "И. Петров",           // инициал и фамилия под „Изготвил“
  "director": {"rank": "КОМИСАР", "name": "Име Фамилия"},   // по подразбиране – директорът от списъка на служителите
  "date": "25.09.2026"
}

Сумите се смятат по формулите на шаблона (ставка = 3,5 % × 620,20 EUR на час, режийни 10 %
от труда, консумативи по единични цени), а общата сума се изписва словом. Жълтото
оцветяване на полетата за попълване се изчиства. .xlsx не изисква нищо допълнително; .xls (само при изрично
искане) изисква LibreOffice (soffice).
"""
import argparse
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile

import openpyxl
from openpyxl.styles import PatternFill

HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = os.path.join(HERE, "..", "assets", "smetka_template.xlsx")

RATE = 0.035 * 620.2          # E11 – the same formula stands in the template; the test compares the two
NUMBERS = ("pages", "copies", "hours", "a4_bw", "a4_color", "a3_bw", "hours_weekend", "hours_holiday", "complex_addon")

# Данните на експерта по подразбиране (_shared/profile.md); другият експерт подава своите.
DEFAULT_TITLE = "инспектор"
DEFAULT_POSITION = "[длъжност на експерта]"
PRICE = {"a4_bw": 0.35, "a4_color": 1.0, "a3_bw": 0.35, "ogled": 1.0}

# ---------- числа словом ----------
_UNITS_N = ["нула", "едно", "две", "три", "четири", "пет", "шест", "седем", "осем", "девет",
            "десет", "единадесет", "дванадесет", "тринадесет", "четиринадесет", "петнадесет",
            "шестнадесет", "седемнадесет", "осемнадесет", "деветнадесет"]
_TENS = ["", "", "двадесет", "тридесет", "четиридесет", "петдесет", "шестдесет", "седемдесет",
         "осемдесет", "деветдесет"]
_HUND = ["", "сто", "двеста", "триста", "четиристотин", "петстотин", "шестстотин", "седемстотин",
         "осемстотин", "деветстотин"]


def _unit(n, gender):
    if n == 1:
        return {"n": "едно", "m": "един", "f": "една"}[gender]
    if n == 2:
        return {"n": "две", "m": "два", "f": "две"}[gender]
    return _UNITS_N[n]


def _parts_under_1000(n, gender):
    parts = []
    h, r = divmod(n, 100)
    if h:
        parts.append(_HUND[h])
    if r:
        if r < 20:
            parts.append(_unit(r, gender))
        else:
            t, u = divmod(r, 10)
            parts.append(_TENS[t])
            if u:
                parts.append(_unit(u, gender))
    return parts


def _join(parts):
    if len(parts) <= 1:
        return "".join(parts)
    return " ".join(parts[:-1]) + " и " + parts[-1]


def words(n, gender="n"):
    """Цяло число 0–999 999 словом. gender: n (евро), m (евроцент), f."""
    if n == 0:
        return "нула"
    th, rest = divmod(n, 1000)
    out = []
    if th:
        out.append("хиляда" if th == 1 else _join(_parts_under_1000(th, "f")) + " хиляди")
    rest_parts = _parts_under_1000(rest, gender)
    if not rest_parts:
        return out[0]
    if not out:
        return _join(rest_parts)
    if len(rest_parts) == 1:
        return out[0] + " и " + rest_parts[0]
    return out[0] + " " + _join(rest_parts)


def amount_words(total):
    cents_total = int(round(total * 100))
    eur, cents = divmod(cents_total, 100)
    s = f"{words(eur, 'n')} евро"
    if cents:
        s += f" и {words(cents, 'm')} {'евроцент' if cents == 1 else 'евроцента'}"
    return s[0].upper() + s[1:] + "."


# ---------- сметка ----------
def apply_defaults(d):
    """Правило на експерта: часове труд = страниците на ЕПТЗ - 2 (напр. 6 стр. → 4 ч.);
    листа А4 = страници × екземпляри. Изрично подадени "hours"/"a4_bw" имат предимство."""
    d = dict(d)
    for k in NUMBERS:                      # numbers may come as text ("4", "4,5") or as null
        v = d.get(k)
        if v is None or v == "":
            d.pop(k, None)
            continue
        try:
            d[k] = float(str(v).replace(",", ".")) if not isinstance(v, (int, float)) else v
        except ValueError:
            sys.exit(f"BAD_INPUT: „{k}“ трябва да е число, а е „{v}“")
        if d[k] < 0:
            sys.exit(f"BAD_INPUT: „{k}“ не може да е отрицателно ({v})")
        if d[k] == int(d[k]):
            d[k] = int(d[k])
    pages = int(d.get("pages") or 0)
    if pages and "hours" not in d:
        d["hours"] = max(pages - 2, 1)
    if pages and "a4_bw" not in d:
        d["a4_bw"] = pages * int(d.get("copies") or 2)
    if not d.get("hours"):
        sys.exit('BAD_INPUT: няма часове труд – подай "pages" (страниците на заключението) или "hours"')
    return d


def staff_director():
    """The director of РДПБЗН – Плевен from the staff list, or None (no key, no answer)."""
    import importlib.util
    path = os.path.join(HERE, "..", "_shared", "scripts", "sluzhiteli.py")
    try:
        spec = importlib.util.spec_from_file_location("sluzhiteli", path)
        mod = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(mod)
        return mod.direktor()
    except Exception:  # noqa: BLE001
        return None


def compute(d):
    labour = d.get("hours", 0) * RATE
    overhead = labour * 0.1
    cons = sum(d.get(k, 0) * PRICE[k] for k in ("a4_bw", "a4_color", "a3_bw"))
    cons += (1 if d.get("ogled") else 0) * PRICE["ogled"]
    addons = d.get("hours_weekend", 0) * RATE * 0.5 + d.get("hours_holiday", 0) * RATE * 1 \
        + d.get("complex_addon", 0)
    return round(labour + overhead + cons + addons, 2)


def fill(d, xlsx_out):
    wb = openpyxl.load_workbook(TEMPLATE)
    ws = wb.active
    dots = lambda v, fallback: v if v else fallback
    title = (d.get("expert_title") or DEFAULT_TITLE).strip()
    ws["A7"] = (
        "по чл.23, ал.3 от Наредба № Н-1/14.02.2023г. за вписването, квалификацията и "
        "възнагражденията на вещите лица за направените разходи за труд, консумативи и режийни "
        f"разноски по експертиза № {dots(d.get('ekspertiza_no'), '………….......................')}"
        f"/{dots(d.get('ekspertiza_year'), '…….…….')} г., назначена от {d.get('naznachil', '……………')} "
        f"при {d.get('pri', '……………')} с постановление за назначаване на експертиза по "
        f"{d.get('delo_vid', 'ДП')} № {d.get('delo_no', '……………')} по описа на "
        f"{d.get('po_opisa_na', '……………')} извършена от {title} {d.get('izvarshena_ot', '……………')} – "
        f"{d.get('expert_position') or DEFAULT_POSITION}"
    )
    ws["D11"] = d.get("hours", 0)
    ws["D16"] = d.get("a4_bw", 0)
    ws["D17"] = d.get("a4_color", 0)
    ws["D18"] = d.get("a3_bw", 0)
    ws["D19"] = 1 if d.get("ogled") else 0
    if d.get("hours_weekend"):
        ws["D23"] = d["hours_weekend"]
    if d.get("hours_holiday"):
        ws["D24"] = d["hours_holiday"]
    if d.get("complex_addon"):
        ws["D25"], ws["E25"] = 1, d["complex_addon"]
    total = compute(d)
    ws["A30"] = f"         Обща стойност (словом): {amount_words(total)}"
    ws["A43"] = f"/ {title} {d.get('izgotvil_short', '[И. Фамилия]')} /"
    dr = d.get("director") or staff_director() or {}
    if not dr.get("name"):
        print("MISSING: director – подай \"director\": {\"rank\", \"name\"} (инструментът staff, dl „директор“)", file=sys.stderr)
    ws["D44"] = (dr.get("rank") or "[ЗВАНИЕ]").upper()
    ws["D45"] = "   " + (dr.get("name") or "[Име Фамилия]")
    if d.get("date"):
        ws["A45"] = f"{d['date']} г." if not str(d["date"]).endswith("г.") else d["date"]
    # изчисти жълтото
    for row in ws.iter_rows():
        for c in row:
            if c.fill is not None and c.fill.fill_type and str(c.fill.fgColor.rgb).upper() == "FFFFFF00":
                c.fill = PatternFill(fill_type=None)
    # типографски правила (_shared/conventions.md): тирето е „–“ (не „-“ и не „—“), „№ 4“
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and not c.value.startswith("="):
                v = re.sub(r"№\s*(?=\S)", "№ ", c.value).replace("—", "–")
                c.value = re.sub(r"(^|\s)-(?=\s)", lambda m: m.group(1) + "–", v)
    wb.save(xlsx_out)
    return total


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("data")
    ap.add_argument("-o", "--out", required=True, help=".xlsx изходен файл (.xls само при изрично искане — изисква LibreOffice)")
    a = ap.parse_args()
    d = apply_defaults(json.load(open(a.data, encoding="utf-8")))
    out = os.path.abspath(a.out)
    if out.lower().endswith(".xlsx"):
        total = fill(d, out)
    else:
        tmp = tempfile.mkdtemp()
        x = os.path.join(tmp, os.path.splitext(os.path.basename(out))[0] + ".xlsx")
        total = fill(d, x)
        soffice = shutil.which("soffice") or shutil.which("libreoffice")
        if not soffice:
            sys.exit("NO_LIBREOFFICE: запиши като .xlsx или инсталирай LibreOffice")
        subprocess.run([soffice, "--headless", "--convert-to", "xls", "--outdir", tmp, x],
                       check=True, capture_output=True)
        shutil.move(os.path.join(tmp, os.path.splitext(os.path.basename(x))[0] + ".xls"), out)
    print(f"✅ {out}  обща сума: {total:.2f} EUR — {amount_words(total)}")


if __name__ == "__main__":
    main()
